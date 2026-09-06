"""The single SahiPack report service and its PDF/DOCX renderers."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PillowImage
from PIL import ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    SimpleDocTemplate,
)

from app.rule_engine.assessment import chart_data, generate_initial_assessment

NAVY = "17324D"
TEAL = "087E8B"
CREAM = "F7F1E6"
GOLD = "D7A928"
RED = "B42318"
AMBER = "B54708"
GREEN = "067647"
GRAY = "667085"
LIGHT = "F2F4F7"
STATUS_COLORS = {
    "COMPLIANT": GREEN,
    "POTENTIAL VIOLATION": RED,
    "NEEDS VERIFICATION": AMBER,
    "NOT APPLICABLE": GRAY,
}


def build_initial_assessment(product: dict[str, Any], findings: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    return generate_initial_assessment(product, findings=findings)


def build_officer_verified_report(inspection: dict[str, Any]) -> dict[str, Any]:
    assessment = inspection.get("assessment") or {}
    findings = inspection.get("findings", [])
    ai_findings = [item for item in findings if item.get("source") == "AI"]
    officer_findings = [item for item in findings if item.get("source") == "OFFICER"]
    comments = inspection.get("comments", [])
    rule_results = assessment.get("rule_results", [])
    return {
        "title": "OFFICER-VERIFIED COMPLIANCE REPORT",
        "stage": "OFFICER_VERIFIED",
        "inspection_summary": {
            "inspection_id": inspection.get("id"),
            "status": inspection.get("status"),
            "final_outcome": inspection.get("final_outcome"),
            "verified_by": inspection.get("verified_by"),
            "verified_at": inspection.get("verified_at"),
            "product": assessment.get("product", {}),
            "screening_score": assessment.get("screening_score"),
            "screening_score_note": "Internal screening indicator - not an official legal determination.",
        },
        "product_declaration_analysis": {"product": assessment.get("product", {}), "declarations": assessment.get("declarations", [])},
        "rule_by_rule_compliance": [
            {
                **result,
                "officer_decision": next((item.get("officer_decision") for item in ai_findings if item.get("applicable_rule") == result.get("rule_number")), None),
                "officer_comment": next((item.get("officer_comment") for item in ai_findings if item.get("applicable_rule") == result.get("rule_number")), None),
            }
            for result in rule_results
        ],
        "visual_packaging_analysis": [item for item in findings if item.get("category") in {"Font & Readability", "Principal Display Panel", "Misleading / Deceptive Packaging", "Quantity & Unit"}],
        "evidence_photographs": inspection.get("evidence", []),
        "officer_verification_record": ai_findings,
        "officer_added_findings": officer_findings,
        "officer_observations": inspection.get("observations", []),
        "overall_inspection_comments": comments,
        "final_decision": {
            "outcome": inspection.get("final_outcome"),
            "confirmed_violations": [item for item in findings if item.get("officer_decision") == "CONFIRM_VIOLATION" or item.get("source") == "OFFICER" and item.get("status") == "POTENTIAL_VIOLATION"],
            "dismissed_concerns": [item for item in findings if item.get("officer_decision") == "DISMISS_CONCERN"],
            "further_inspection": [item for item in findings if item.get("officer_decision") == "NEEDS_FURTHER_INSPECTION"],
            "officer_added_findings": officer_findings,
            "observations": inspection.get("observations", []),
            "comments": comments,
        },
        "charts": chart_data(rule_results, findings, assessment.get("screening_score", 0)),
        "ai_assessment_preserved": assessment,
    }


def _text(value: Any, fallback: str = "Not available") -> str:
    if value is None or value == "":
        return fallback
    if isinstance(value, (datetime, date)):
        return value.strftime("%d %b %Y, %H:%M UTC") if isinstance(value, datetime) else value.strftime("%d %b %Y")
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, dict):
        for key in ("value", "name", "title", "description", "text", "id"):
            if value.get(key) not in (None, ""):
                return _text(value[key], fallback)
        return "; ".join(f"{key.replace('_', ' ').title()}: {_text(item)}" for key, item in value.items() if item not in (None, "", [], {})) or fallback
    if isinstance(value, list):
        return "; ".join(_text(item) for item in value) if value else fallback
    return str(value)


def _status(value: Any) -> str:
    return _text(value).replace("_", " ").upper()


def _confidence(value: Any) -> str:
    if value in (None, ""):
        return "Not available"
    try:
        number = float(value)
        return f"{number:.0%}" if number <= 1 else f"{number:.0f}%"
    except (TypeError, ValueError):
        return _text(value)


def _evidence_text(evidence: Any) -> str:
    if not evidence:
        return "No evidence reference"
    if isinstance(evidence, list):
        return "; ".join(_evidence_text(item) for item in evidence)
    if isinstance(evidence, dict):
        parts = []
        for key in ("evidence_id", "image_id", "file_path", "description", "caption", "source_image"):
            if evidence.get(key) not in (None, ""):
                parts.append(f"{key.replace('_', ' ').title()}: {_text(evidence[key])}")
        return "; ".join(parts) or "Evidence available"
    return _text(evidence)


def _image_path(item: Any) -> Path | None:
    if not isinstance(item, dict):
        return None
    candidate = item.get("file_path") or item.get("image_path") or item.get("path")
    path = Path(candidate) if candidate else None
    return path if path and path.is_file() else None


def _product_rows(product: dict[str, Any]) -> list[list[str]]:
    labels = (("product_name", "Product Name"), ("name", "Product Name"), ("brand", "Brand"), ("category", "Category"), ("manufacturer_name", "Manufacturer"), ("manufacturer_address", "Manufacturer Address"), ("packer_name", "Packer"), ("importer_name", "Importer"), ("net_quantity", "Net Quantity"), ("mrp", "MRP"), ("month_year_of_manufacture", "Manufacture / Packing Date"))
    rows, seen = [], set()
    for key, label in labels:
        if key in product and label not in seen:
            rows.append([label, _text(product[key])])
            seen.add(label)
    return rows or [["Product information", "Not available"]]


def _declaration_rows(declarations: Iterable[dict[str, Any]]) -> list[list[str]]:
    rows = []
    for item in declarations:
        rows.append([_text(item.get("declaration")), _text(item.get("detected_value")), _status(item.get("status")), _text(item.get("rule")), _evidence_text(item.get("evidence"))])
    return rows or [["No declarations available", "-", "-", "-", "-"]]


def _finding_rows(findings: Iterable[dict[str, Any]]) -> list[list[str]]:
    rows = []
    for item in findings:
        rows.append([_text(item.get("applicable_rule")), _text(item.get("category")), _text(item.get("title") or item.get("description")), _status(item.get("original_ai_status") or item.get("status")), _confidence(item.get("ai_confidence") or item.get("confidence")), _status(item.get("officer_decision")), _text(item.get("officer_comment")), _evidence_text(item.get("evidence") or item.get("ai_evidence"))])
    return rows or [["-", "-", "No findings available", "-", "-", "-", "-", "-"]]


def _compact_finding_rows(findings: Iterable[dict[str, Any]]) -> list[list[str]]:
    rows = []
    for item in findings:
        rows.append([_text(item.get("title") or item.get("description")), _text(item.get("applicable_rule")), _status(item.get("original_ai_status") or item.get("status")), _confidence(item.get("ai_confidence") or item.get("confidence")), _status(item.get("officer_decision")), _evidence_text(item.get("ai_evidence") or item.get("evidence"))])
    return rows or [["No findings available", "-", "-", "-", "-", "-"]]


def _verification_rows(findings: Iterable[dict[str, Any]]) -> list[list[str]]:
    rows = []
    for item in findings:
        rows.append([_text(item.get("title") or item.get("description")), _status(item.get("original_ai_status")), _confidence(item.get("ai_confidence")), _status(item.get("officer_decision")), _text(item.get("officer_comment")), f"{_text(item.get('verified_by'))} / {_text(item.get('verified_at'))}"])
    return rows or [["No AI findings", "-", "-", "-", "-", "-"]]


def _chart_rows(report: dict[str, Any]) -> list[tuple[str, dict[str, Any]]]:
    charts = report.get("charts") or {}
    return [("Compliance Distribution", charts.get("compliance_distribution", {})), ("Findings by Category", charts.get("findings_by_category", {})), ("Screening Score Breakdown", charts.get("screening_score_breakdown", {})), ("Finding Resolution", charts.get("finding_resolution", {}))]


def _chart_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
    for candidate in candidates:
        if candidate.is_file():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def _chart_png(title: str, values: dict[str, Any], width: int = 1200, height: int = 620) -> bytes:
    image = PillowImage.new("RGB", (width, height), "#F7F1E6")
    draw = ImageDraw.Draw(image)
    title_font = _chart_font(34, bold=True)
    label_font = _chart_font(23)
    value_font = _chart_font(25, bold=True)
    small_font = _chart_font(18)
    draw.text((46, 32), title.upper(), fill="#17324D", font=title_font)
    pairs = [(str(key).replace("_", " "), float(value)) for key, value in values.items() if isinstance(value, (int, float))]
    if not pairs or sum(value for _, value in pairs) == 0:
        draw.text((46, 285), "No data available", fill="#667085", font=label_font)
    elif title == "Compliance Distribution":
        total = sum(value for _, value in pairs)
        center = (300, 350)
        radius = 185
        start = 0.0
        palette = ["#067647", "#B42318", "#B54708", "#667085"]
        for index, (_, value) in enumerate(pairs):
            extent = value / total * 360
            draw.pieslice((center[0] - radius, center[1] - radius, center[0] + radius, center[1] + radius), start, start + extent, fill=palette[index % len(palette)])
            start += extent
        draw.ellipse((center[0] - 86, center[1] - 86, center[0] + 86, center[1] + 86), fill="#F7F1E6")
        draw.text((center[0] - 32, center[1] - 30), str(int(total)), fill="#17324D", font=_chart_font(42, bold=True))
        draw.text((center[0] - 44, center[1] + 20), "checks", fill="#667085", font=small_font)
        for index, (label, value) in enumerate(pairs):
            y = 150 + index * 76
            draw.rounded_rectangle((610, y, 638, y + 28), radius=6, fill=palette[index % len(palette)])
            draw.text((660, y - 2), label.title(), fill="#344054", font=label_font)
            draw.text((1030, y - 2), str(int(value) if value.is_integer() else round(value, 2)), fill="#17324D", font=value_font)
    elif title == "Screening Score Breakdown":
        score = pairs[0][1]
        draw.text((52, 150), "AI SCREENING SCORE", fill="#087E8B", font=label_font)
        draw.text((48, 205), f"{score:.2f}", fill="#17324D", font=_chart_font(92, bold=True))
        draw.text((52, 320), "Internal screening metric - not a legal, government,", fill="#667085", font=small_font)
        draw.text((52, 350), "or official compliance score.", fill="#667085", font=small_font)
        draw.rounded_rectangle((52, 455, width - 52, 493), radius=19, fill="#E4E7EC")
        progress = max(0, min(score, 100)) / 100
        draw.rounded_rectangle((52, 455, 52 + int((width - 104) * progress), 493), radius=19, fill="#D7A928")
        draw.text((52, 515), "0", fill="#667085", font=small_font)
        draw.text((width - 78, 515), "100", fill="#667085", font=small_font)
    else:
        maximum = max(value for _, value in pairs) or 1
        left = 300
        right = width - 90
        bar_height = 42
        gap = 75 if len(pairs) <= 4 else 55
        for index, (label, value) in enumerate(pairs):
            y = 145 + index * gap
            draw.text((48, y + 4), label.title()[:25], fill="#344054", font=label_font)
            draw.rounded_rectangle((left, y, right, y + bar_height), radius=12, fill="#E4E7EC")
            bar_end = left + int((right - left) * value / maximum)
            draw.rounded_rectangle((left, y, max(left + 12, bar_end), y + bar_height), radius=12, fill="#087E8B" if title == "Findings by Category" else "#D7A928")
            draw.text((min(bar_end + 16, right - 20), y + 4), str(int(value) if value.is_integer() else round(value, 2)), fill="#17324D", font=value_font)
    output = BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _pdf_chart_grid(chart_sets: list[tuple[str, dict[str, Any]]], styles: dict[str, ParagraphStyle]) -> Table:
    cells = []
    for title, values in chart_sets:
        card = Table([[Image(BytesIO(_chart_png(title, values)), width=82 * mm, height=42 * mm)]], colWidths=[82 * mm], rowHeights=[42 * mm])
        card.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{CREAM}")), ("BOX", (0, 0), (-1, -1), 0.45, colors.HexColor("#D0D5DD")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
        cells.append(card)
    while len(cells) % 2:
        cells.append(Spacer(82 * mm, 42 * mm))
    grid = Table([cells[index:index + 2] for index in range(0, len(cells), 2)], colWidths=[85 * mm, 85 * mm], hAlign="LEFT")
    grid.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
    return grid


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("ReportTitle", parent=base["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=colors.HexColor(f"#{NAVY}"), alignment=TA_CENTER, spaceAfter=8),
        "hero": ParagraphStyle("ReportHero", parent=base["Title"], fontName="Helvetica-Bold", fontSize=28, leading=31, textColor=colors.HexColor(f"#{NAVY}"), alignment=TA_LEFT, spaceAfter=12),
        "subtitle": ParagraphStyle("ReportSubtitle", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=13, leading=16, textColor=colors.HexColor(f"#{TEAL}"), alignment=TA_CENTER, spaceAfter=18),
        "eyebrow": ParagraphStyle("ReportEyebrow", parent=base["BodyText"], fontName="Helvetica-Bold", fontSize=7, leading=9, textColor=colors.HexColor(f"#{TEAL}"), spaceAfter=8),
        "heading": ParagraphStyle("SectionHeading", parent=base["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=20, textColor=colors.HexColor(f"#{NAVY}"), spaceBefore=4, spaceAfter=10),
        "subheading": ParagraphStyle("SubHeading", parent=base["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=colors.HexColor(f"#{TEAL}"), spaceBefore=8, spaceAfter=6),
        "body": ParagraphStyle("ReportBody", parent=base["BodyText"], fontName="Helvetica", fontSize=8.5, leading=11, textColor=colors.HexColor("#344054")),
        "small": ParagraphStyle("ReportSmall", parent=base["BodyText"], fontName="Helvetica", fontSize=7, leading=9, textColor=colors.HexColor("#344054")),
    }


def _p(value: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(_text(value).replace("&", "&amp;").replace("<", "&lt;"), style)


def _pdf_table(headers: list[str], rows: list[list[Any]], styles: dict[str, ParagraphStyle], widths: list[float] | None = None) -> Table:
    data = [[_p(header, styles["small"]) for header in headers]]
    data.extend([[_p(cell, styles["small"]) for cell in row] for row in rows])
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.HexColor(f"#{GOLD}")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{CREAM}")]), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    return table


def _pdf_metadata_strip(fields: list[tuple[str, Any]], styles: dict[str, ParagraphStyle]) -> Table:
    labels = [_p(label.upper(), styles["small"]) for label, _ in fields]
    values = [_p(value, styles["body"]) for _, value in fields]
    table = Table([labels, values], colWidths=[173 * mm / len(fields)] * len(fields), hAlign="LEFT")
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{CREAM}")), ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.HexColor(f"#{GOLD}")), ("LINEBELOW", (0, 1), (-1, 1), 0.25, colors.HexColor("#D0D5DD")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7), ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8)]))
    return table


def _pdf_outcome_block(outcome: Any, styles: dict[str, ParagraphStyle]) -> Table:
    table = Table([[Paragraph("FINAL INSPECTION OUTCOME", styles["small"])], [Paragraph(_status(outcome), styles["hero"])]], colWidths=[173 * mm])
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor(f"#{NAVY}")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor(f"#{GOLD}")), ("TEXTCOLOR", (0, 1), (-1, 1), colors.HexColor(f"#{GOLD}")), ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor(f"#{GOLD}")), ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14), ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
    return table


def _pdf_finding_card(item: dict[str, Any], styles: dict[str, ParagraphStyle], officer: bool = False) -> KeepTogether:
    status = _status(item.get("status") or item.get("original_ai_status"))
    status_color = colors.HexColor(f"#{STATUS_COLORS.get(status, GRAY)}")
    title = "OFFICER-ADDED FINDING" if officer else _text(item.get("title") or item.get("description"))
    rows = [
        [_p(title, styles["subheading"]), _p(status, styles["small"])],
        [_p("Description", styles["small"]), _p(item.get("description"), styles["body"])],
        [_p("Applicable Rule", styles["small"]), _p(item.get("applicable_rule"), styles["body"])],
        [_p("AI Detection Confidence", styles["small"]), _p(_confidence(item.get("ai_confidence")), styles["body"])],
        [_p("Original AI Status", styles["small"]), _p(_status(item.get("original_ai_status")), styles["body"])],
        [_p("AI Evidence", styles["small"]), _p(_evidence_text(item.get("ai_evidence") or item.get("evidence")), styles["body"])],
        [_p("Officer Decision", styles["small"]), _p(_status(item.get("officer_decision")), styles["body"])],
        [_p("Officer Comment", styles["small"]), _p(item.get("officer_comment"), styles["body"])],
        [_p("Verified By / Time", styles["small"]), _p(f"{_text(item.get('verified_by'))} / {_text(item.get('verified_at'))}", styles["body"])],
    ]
    card = Table(rows, colWidths=[42 * mm, 131 * mm], hAlign="LEFT")
    card.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{LIGHT}")), ("BACKGROUND", (1, 0), (1, 0), status_color), ("TEXTCOLOR", (1, 0), (1, 0), colors.white), ("BOX", (0, 0), (-1, -1), 0.7, status_color), ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D0D5DD")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]))
    return KeepTogether([card, Spacer(1, 8)])


def _pdf_footer(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont("Helvetica-Bold", 7)
    canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
    canvas.drawString(18 * mm, A4[1] - 11 * mm, "SAHIPACK | OFFICER-VERIFIED COMPLIANCE REPORT")
    canvas.setStrokeColor(colors.HexColor("#D0D5DD"))
    canvas.line(18 * mm, 14 * mm, A4[0] - 18 * mm, 14 * mm)
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.drawString(18 * mm, 9 * mm, "SahiPack | Officer-verified inspection record")
    canvas.drawRightString(A4[0] - 18 * mm, 9 * mm, f"Page {document.page}")
    canvas.restoreState()


def _pdf_section(report: dict[str, Any], styles: dict[str, ParagraphStyle], title: str, content: list[Any]) -> list[Any]:
    return [Paragraph(title, styles["heading"]), *content, PageBreak()]


def render_pdf(report: dict[str, Any]) -> bytes:
    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm, topMargin=18 * mm, bottomMargin=20 * mm, title=report["title"])
    styles = _pdf_styles()
    summary = report.get("inspection_summary", {})
    product = summary.get("product") or {}
    story: list[Any] = [Paragraph("SAHIPACK", styles["eyebrow"]), Paragraph("LEGAL METROLOGY • EVIDENCE-LED INSPECTION", styles["eyebrow"]), Paragraph("OFFICER-VERIFIED\nCOMPLIANCE\nREPORT.", styles["hero"]), Paragraph("An evidence-backed inspection record containing product declarations, applicable rules, AI screening findings and officer verification.", styles["body"]), Spacer(1, 16), _pdf_outcome_block(summary.get("final_outcome"), styles), Spacer(1, 15), _pdf_metadata_strip([["Inspection", _text(summary.get("inspection_id"))], ["Product", _text(product.get("product_name") or product.get("name"))], ["Brand", _text(product.get("brand"))], ["Category", _text(product.get("category"))], ["Net Quantity", _text(product.get("net_quantity"))]], styles), Spacer(1, 12), _p("AI findings are screening results and require officer verification.", styles["body"]), Spacer(1, 12)]
    story.append(_pdf_metadata_strip([["Inspection Status", _status(summary.get("status"))], ["Verified By", _text(summary.get("verified_by"))], ["Verification Time", _text(summary.get("verified_at"))], ["Screening Score", _text(summary.get("screening_score"))]], styles))
    chart_sets = _chart_rows(report)
    distribution = (report.get("charts") or {}).get("compliance_distribution", {})
    stats = [["Applicable checks", sum(value for value in distribution.values() if isinstance(value, (int, float)))], ["Compliant", distribution.get("COMPLIANT", 0)], ["Potential violations", distribution.get("POTENTIAL_VIOLATION", 0)], ["Needs verification", distribution.get("NEEDS_VERIFICATION", 0)]]
    story.extend(_pdf_section(report, styles, "INITIAL COMPLIANCE ASSESSMENT", [Paragraph("INITIAL ASSESSMENT", styles["eyebrow"]), Paragraph("See what the screening found.", styles["hero"]), Paragraph("SahiPack identifies declarations, surfaces potential concerns and records uncertainty for officer review. This assessment is preliminary and is not a legal determination.", styles["body"]), Spacer(1, 8), _pdf_table(["Summary", "Count"], stats, styles, [125 * mm, 48 * mm]), Spacer(1, 8), _pdf_chart_grid(chart_sets, styles), Spacer(1, 5), Paragraph("PRELIMINARY AI ASSESSMENT | OFFICER REVIEW REQUIRED", styles["subheading"])]))
    declarations = report.get("product_declaration_analysis", {})
    story.extend(_pdf_section(report, styles, "PRODUCT & DECLARATION ANALYSIS", [Paragraph("PRODUCT INFORMATION", styles["subheading"]), _pdf_table(["Field", "Value"], _product_rows(declarations.get("product", {})), styles, [55 * mm, 118 * mm]), Spacer(1, 12), Paragraph("DECLARATION ANALYSIS", styles["subheading"]), _pdf_table(["Declaration / Requirement", "Detected Value", "Status", "Rule", "Evidence"], _declaration_rows(declarations.get("declarations", [])), styles, [37 * mm, 35 * mm, 30 * mm, 22 * mm, 49 * mm])]))
    rule_content = [Paragraph("AI findings and officer decisions are shown together without overwriting the original assessment.", styles["body"]), _pdf_table(["Finding", "Rule", "AI Status", "Confidence", "Officer Decision", "Evidence"], _compact_finding_rows(report.get("rule_by_rule_compliance", [])), styles, [47 * mm, 18 * mm, 28 * mm, 20 * mm, 30 * mm, 30 * mm])]
    commented_findings = [[_text(item.get("title") or item.get("description")), _text(item.get("officer_comment"))] for item in report.get("rule_by_rule_compliance", []) if item.get("officer_comment")]
    if commented_findings:
        rule_content.extend([Paragraph("OFFICER COMMENTS", styles["subheading"]), _pdf_table(["Finding", "Comment"], commented_findings, styles, [55 * mm, 118 * mm])])
    story.extend(_pdf_section(report, styles, "RULE-BY-RULE COMPLIANCE", rule_content))
    visual = report.get("visual_packaging_analysis", [])
    physical = (report.get("ai_assessment_preserved") or {}).get("physical_quantity")
    story.extend(_pdf_section(report, styles, "VISUAL & PACKAGING ANALYSIS", [_pdf_table(["Category", "Finding", "Rule", "Status", "AI Confidence", "Officer Comment", "Evidence"], [[_text(item.get("category")), _text(item.get("title") or item.get("description")), _text(item.get("applicable_rule")), _status(item.get("status") or item.get("original_ai_status")), _confidence(item.get("ai_confidence")), _text(item.get("officer_comment")), _evidence_text(item.get("evidence") or item.get("ai_evidence"))] for item in visual] or [["-", "No visual findings available", "-", "-", "-", "-", "-"]], styles, [22 * mm, 32 * mm, 15 * mm, 22 * mm, 18 * mm, 34 * mm, 31 * mm]), Spacer(1, 12), Paragraph("QUANTITY SAFEGUARD", styles["subheading"]), _p((physical or {}).get("message") or "Physical quantity not verified from image evidence.", styles["body"])]))
    evidence = report.get("evidence_photographs", [])
    evidence_rows = [[_text(item.get("evidence_id") or item.get("image_id")), _text(item.get("finding_id") or item.get("related_finding")), _text(item.get("rule") or item.get("applicable_rule")), _text(item.get("caption") or item.get("description")), _text(item.get("source"))] for item in evidence if isinstance(item, dict)] or [["-", "-", "-", "No photographic evidence available.", "-"]]
    evidence_content: list[Any] = [_pdf_table(["Evidence ID", "Related Finding", "Rule", "Description / Caption", "Source"], evidence_rows, styles, [25 * mm, 32 * mm, 22 * mm, 68 * mm, 25 * mm])]
    for item in evidence:
        path = _image_path(item)
        if path:
            evidence_content.extend([Spacer(1, 8), Paragraph(_text(item.get("caption") or item.get("description")), styles["small"]), Image(str(path), width=70 * mm, height=50 * mm)])
    story.extend(_pdf_section(report, styles, "EVIDENCE & PHOTOGRAPHS", evidence_content))
    verification = report.get("officer_verification_record", [])
    verification_content = [Paragraph("AI DETECTION → OFFICER VERIFICATION", styles["subheading"]), Paragraph("Original AI status remains distinct from the officer decision.", styles["body"]), _pdf_table(["Finding", "AI Status", "Confidence", "Officer Decision", "Officer Comment", "Verified By / Time"], _verification_rows(verification), styles, [36 * mm, 25 * mm, 19 * mm, 28 * mm, 42 * mm, 23 * mm]), Paragraph("OVERALL INSPECTION COMMENTS", styles["subheading"]), _pdf_table(["Officer", "Timestamp", "Comment"], [[_text(item.get("officer")), _text(item.get("timestamp")), _text(item.get("comment"))] for item in report.get("overall_inspection_comments", [])] or [["-", "-", "No overall inspection comments."]], styles, [25 * mm, 38 * mm, 131 * mm])]
    story.extend(_pdf_section(report, styles, "OFFICER VERIFICATION RECORD", [Paragraph("Human Verification: VERIFIED", styles["subheading"]), *verification_content]))
    officer_findings = report.get("officer_added_findings", [])
    observations = report.get("officer_observations", [])
    officer_content: list[Any] = [Paragraph("OFFICER-ADDED | SOURCE: OFFICER", styles["subheading"]), _pdf_table(["Finding", "Rule", "Category", "Status", "Officer Comment", "Evidence"], [[_text(item.get("title") or item.get("description")), _text(item.get("applicable_rule")), _text(item.get("category")), _status(item.get("status")), _text(item.get("officer_comment")), _evidence_text(item.get("evidence"))] for item in officer_findings] or [["No officer-added findings", "-", "-", "-", "-", "-"]], styles, [38 * mm, 18 * mm, 27 * mm, 25 * mm, 40 * mm, 25 * mm]), Spacer(1, 8), Paragraph("OFFICER OBSERVATIONS", styles["subheading"]), _pdf_table(["Observation", "Evidence", "Officer", "Timestamp"], [[_text(item.get("description")), _evidence_text(item.get("evidence")), _text(item.get("officer")), _text(item.get("timestamp"))] for item in observations] or [["No officer observations.", "-", "-", "-"]], styles, [70 * mm, 48 * mm, 25 * mm, 50 * mm])]
    story.extend(_pdf_section(report, styles, "OFFICER-ADDED FINDINGS & OBSERVATIONS", officer_content))
    decision = report.get("final_decision", {})
    story.extend(_pdf_section(report, styles, "FINAL DECISION", [Paragraph(f"FINAL OUTCOME: {_status(decision.get('outcome'))}", styles["subtitle"]), _pdf_table(["Decision Summary", "Count"], [["Confirmed violations", len(decision.get("confirmed_violations", []))], ["Dismissed concerns", len(decision.get("dismissed_concerns", []))], ["Needs further inspection", len(decision.get("further_inspection", []))], ["Officer-added findings", len(decision.get("officer_added_findings", []))], ["Officer observations", len(decision.get("observations", []))]], styles, [125 * mm, 48 * mm]), Spacer(1, 12), Paragraph("Human Verification: VERIFIED", styles["body"])]))
    if story and isinstance(story[-1], PageBreak):
        story.pop()
    document.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return output.getvalue()


def _docx_cell(cell: Any, value: Any, bold: bool = False) -> None:
    cell.text = _text(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.size = Pt(8)
            run.bold = bold


def _docx_table(document: Document, headers: list[str], rows: list[list[Any]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        _docx_cell(cell, header, bold=True)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor.from_string("FFFFFF")
        cell._tc.get_or_add_tcPr().append(__import__("docx").oxml.parse_xml('<w:shd {} w:fill="17324D"/>'.format(__import__("docx").oxml.ns.nsdecls("w"))))
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            _docx_cell(cell, value)


def _docx_finding_card(document: Document, item: dict[str, Any], officer: bool = False) -> None:
    status = _status(item.get("status") or item.get("original_ai_status"))
    title = "OFFICER-ADDED FINDING" if officer else _text(item.get("title") or item.get("description"))
    document.add_heading(title, 2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Status", status),
        ("Description", _text(item.get("description"))),
        ("Applicable Rule", _text(item.get("applicable_rule"))),
        ("AI Detection Confidence", _confidence(item.get("ai_confidence"))),
        ("Original AI Status", _status(item.get("original_ai_status"))),
        ("AI Evidence", _evidence_text(item.get("ai_evidence") or item.get("evidence"))),
        ("Officer Decision", _status(item.get("officer_decision"))),
        ("Officer Comment", _text(item.get("officer_comment"))),
        ("Verified By / Time", f"{_text(item.get('verified_by'))} / {_text(item.get('verified_at'))}"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        _docx_cell(cells[0], label, bold=True)
        _docx_cell(cells[1], value)
    document.add_paragraph()


def _docx_metadata_strip(document: Document, fields: list[tuple[str, Any]]) -> None:
    table = document.add_table(rows=2, cols=len(fields))
    table.style = "Table Grid"
    for index, (label, value) in enumerate(fields):
        _docx_cell(table.cell(0, index), label.upper(), bold=True)
        _docx_cell(table.cell(1, index), value)


def _docx_outcome_block(document: Document, outcome: Any) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14)
    paragraph.paragraph_format.space_after = Pt(14)
    run = paragraph.add_run(f"FINAL INSPECTION OUTCOME\n{_status(outcome)}")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(GOLD)


def _docx_headers_and_footers(document: Document, inspection_id: Any) -> None:
    for section in document.sections:
        header = section.header.paragraphs[0]
        header.text = "SAHIPACK  |  OFFICER-VERIFIED COMPLIANCE REPORT"
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in header.runs:
            run.bold = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(NAVY)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run(f"SAHIPACK  |  Inspection ID: {_text(inspection_id)}  |  Page ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
        for run in footer.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string(GRAY)


def _docx_heading(document: Document, title: str, level: int = 1) -> None:
    heading = document.add_heading(title, level=level)
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)


def _docx_chart(document: Document, title: str, values: dict[str, Any]) -> None:
    document.add_picture(BytesIO(_chart_png(title, values)), width=Inches(6.5))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _docx_chart_grid(document: Document, chart_sets: list[tuple[str, dict[str, Any]]]) -> None:
    table = document.add_table(rows=(len(chart_sets) + 1) // 2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (title, values) in enumerate(chart_sets):
        cell = table.cell(index // 2, index % 2)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run()
        run.add_picture(BytesIO(_chart_png(title, values)), width=Inches(3.15))


def render_docx(report: dict[str, Any]) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)
    summary = report.get("inspection_summary", {})
    product = summary.get("product") or {}
    document.add_heading("SAHIPACK", 0)
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("LEGAL METROLOGY • EVIDENCE-LED INSPECTION")
    document.add_heading("OFFICER-VERIFIED\nCOMPLIANCE\nREPORT.", 1)
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("An evidence-backed inspection record containing product declarations, applicable rules, AI screening findings and officer verification.")
    _docx_outcome_block(document, summary.get("final_outcome"))
    _docx_metadata_strip(document, [("Inspection", _text(summary.get("inspection_id"))), ("Product", _text(product.get("product_name") or product.get("name"))), ("Brand", _text(product.get("brand"))), ("Category", _text(product.get("category"))), ("Net Quantity", _text(product.get("net_quantity")))])
    document.add_paragraph("AI findings are screening results and require officer verification.")
    _docx_metadata_strip(document, [("Status", _status(summary.get("status"))), ("Verified By", _text(summary.get("verified_by"))), ("Verified At", _text(summary.get("verified_at"))), ("Screening Score", _text(summary.get("screening_score")))])

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "INITIAL COMPLIANCE ASSESSMENT")
    document.add_paragraph("INITIAL ASSESSMENT")
    document.add_heading("See what the screening found.", 1)
    document.add_paragraph("SahiPack identifies declarations, surfaces potential concerns and records uncertainty for officer review. This assessment is preliminary and is not a legal determination.")
    distribution = (report.get("charts") or {}).get("compliance_distribution", {})
    _docx_table(document, ["Summary", "Count"], [["Applicable checks", sum(value for value in distribution.values() if isinstance(value, (int, float)))], ["Compliant", distribution.get("COMPLIANT", 0)], ["Potential violations", distribution.get("POTENTIAL_VIOLATION", 0)], ["Needs verification", distribution.get("NEEDS_VERIFICATION", 0)]])
    document.add_paragraph(summary.get("screening_score_note", ""))
    _docx_chart_grid(document, _chart_rows(report))

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "PRODUCT & DECLARATION ANALYSIS")
    _docx_heading(document, "Product Information", 2)
    _docx_table(document, ["Field", "Value"], _product_rows((report.get("product_declaration_analysis") or {}).get("product", {})))
    _docx_heading(document, "Declaration Analysis", 2)
    _docx_table(document, ["Declaration / Requirement", "Detected Value", "Status", "Rule", "Evidence"], _declaration_rows((report.get("product_declaration_analysis") or {}).get("declarations", [])))

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "RULE-BY-RULE COMPLIANCE")
    document.add_paragraph("AI findings and officer decisions are shown together without overwriting the original assessment.")
    _docx_table(document, ["Finding", "Rule", "AI Status", "Confidence", "Officer Decision", "Evidence"], _compact_finding_rows(report.get("rule_by_rule_compliance", [])))
    commented_findings = [[_text(item.get("title") or item.get("description")), _text(item.get("officer_comment"))] for item in report.get("rule_by_rule_compliance", []) if item.get("officer_comment")]
    if commented_findings:
        document.add_heading("Officer Comments", 2)
        _docx_table(document, ["Finding", "Comment"], commented_findings)

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "VISUAL & PACKAGING ANALYSIS")
    visual = report.get("visual_packaging_analysis", [])
    _docx_table(document, ["Category", "Finding", "Rule", "Status", "AI Confidence", "Officer Comment", "Evidence"], [[_text(item.get("category")), _text(item.get("title") or item.get("description")), _text(item.get("applicable_rule")), _status(item.get("status") or item.get("original_ai_status")), _confidence(item.get("ai_confidence")), _text(item.get("officer_comment")), _evidence_text(item.get("evidence") or item.get("ai_evidence"))] for item in visual] or [["-", "No visual findings available", "-", "-", "-", "-", "-"]])
    document.add_heading("Quantity Safeguard", 2)
    document.add_paragraph(_text(((report.get("ai_assessment_preserved") or {}).get("physical_quantity") or {}).get("message") or "Physical quantity not verified from image evidence."))

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "EVIDENCE & PHOTOGRAPHS")
    evidence = report.get("evidence_photographs", [])
    _docx_table(document, ["Evidence ID", "Related Finding", "Rule", "Description / Caption", "Source"], [[_text(item.get("evidence_id") or item.get("image_id")), _text(item.get("finding_id") or item.get("related_finding")), _text(item.get("rule") or item.get("applicable_rule")), _text(item.get("caption") or item.get("description")), _text(item.get("source"))] for item in evidence if isinstance(item, dict)] or [["-", "-", "-", "No photographic evidence available.", "-"]])
    for item in evidence:
        path = _image_path(item)
        if path:
            document.add_paragraph(_text(item.get("caption") or item.get("description")))
            document.add_picture(str(path), width=Inches(3.2))

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "OFFICER VERIFICATION RECORD")
    document.add_paragraph("Human Verification: VERIFIED")
    document.add_heading("AI Detection → Officer Verification", 2)
    document.add_paragraph("Original AI status remains distinct from the officer decision.")
    verification = report.get("officer_verification_record", [])
    _docx_table(document, ["Finding", "AI Status", "Confidence", "Officer Decision", "Officer Comment", "Verified By / Time"], _verification_rows(verification))
    _docx_heading(document, "Overall Inspection Comments", 2)
    _docx_table(document, ["Officer", "Timestamp", "Comment"], [[_text(item.get("officer")), _text(item.get("timestamp")), _text(item.get("comment"))] for item in report.get("overall_inspection_comments", [])] or [["-", "-", "No overall inspection comments."]])

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "OFFICER-ADDED FINDINGS & OBSERVATIONS")
    document.add_heading("Officer-Added | Source: Officer", 2)
    officer_findings = report.get("officer_added_findings", [])
    _docx_table(document, ["Finding", "Rule", "Category", "Status", "Officer Comment", "Evidence"], [[_text(item.get("title") or item.get("description")), _text(item.get("applicable_rule")), _text(item.get("category")), _status(item.get("status")), _text(item.get("officer_comment")), _evidence_text(item.get("evidence"))] for item in officer_findings] or [["No officer-added findings", "-", "-", "-", "-", "-"]])
    document.add_heading("Officer Observations", 2)
    observations = report.get("officer_observations", [])
    _docx_table(document, ["Observation", "Evidence", "Officer", "Timestamp"], [[_text(item.get("description")), _evidence_text(item.get("evidence")), _text(item.get("officer")), _text(item.get("timestamp"))] for item in observations] or [["No officer observations.", "-", "-", "-"]])

    document.add_section(WD_SECTION.NEW_PAGE)
    _docx_heading(document, "FINAL DECISION")
    decision = report.get("final_decision", {})
    document.add_heading(f"FINAL OUTCOME: {_status(decision.get('outcome'))}", 1)
    _docx_table(document, ["Decision Summary", "Count"], [["Confirmed violations", len(decision.get("confirmed_violations", []))], ["Dismissed concerns", len(decision.get("dismissed_concerns", []))], ["Needs further inspection", len(decision.get("further_inspection", []))], ["Officer-added findings", len(decision.get("officer_added_findings", []))], ["Officer observations", len(decision.get("observations", []))]])
    document.add_paragraph("Human Verification: VERIFIED")
    _docx_headers_and_footers(document, summary.get("inspection_id"))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
